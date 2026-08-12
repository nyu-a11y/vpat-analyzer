#!/usr/bin/env python3
"""Generate deterministic, synthetic real-format fixtures for the Stage 0 gate.

All content comes from the committed Northstar JSON model. No production or user
document is read. The output files are safe to commit and intentionally exercise
positive and negative ingestion paths.
"""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import shutil
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont
from pypdf import PdfReader, PdfWriter
from pypdf.generic import ArrayObject, ByteStringObject
from reportlab.lib.colors import Color, HexColor, white
from reportlab.lib.pagesizes import landscape, letter
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfgen import canvas


ROOT = Path(__file__).resolve().parents[2]
SOURCE_PATH = ROOT / "tests/fixtures/synthetic/google-docs-candidate-tables.v1.json"
CATALOG_PATH = ROOT / "config/vpat-2.5-wcag-criteria.v1.json"
OUTPUT_DIR = ROOT / "tests/fixtures/synthetic/real-format"
FIXED_TIME = datetime(2026, 8, 12, 12, 0, 0, tzinfo=timezone.utc)
FIXTURE_PASSWORD = "northstar-fixture"
GENERATOR_VERSION = "1.3.0"
PARSER_DESIGN_PRESET = "google_docs_default"
UNION_CROSS_PAGE_SC = "4.1.1"
DEFAULT_DOCX_MAX_XML_NODES = 200_000
DEFAULT_DOCX_MAX_SOURCE_BYTES = 25 * 1024 * 1024

PURPLE = "#57068C"
INK = "#1F2937"
LIGHT_PURPLE = "#F4EFF8"
LIGHT_GRAY = "#F3F4F6"


def official_vpat_label_for_criterion(criterion: dict, row_index: int) -> str:
    """Shape one catalog item like an official VPAT 2.5Rev criterion cell."""
    active = criterion["activeIn"]
    if criterion["sc"] == "4.1.1":
        suffix = " (Level A) WCAG 2.0 and 2.1 – Always answer ‘Supports’ WCAG 2.2 (obsolete and removed) - Does not apply"
    elif active == ["2.2"]:
        suffix = f" (Level {criterion['level']} 2.2 only)"
    elif active == ["2.1", "2.2"]:
        suffix = f" (Level {criterion['level']} 2.1 and 2.2)"
    else:
        suffix = f" (Level {criterion['level']})"
    shaped = f"{criterion['sc']} {criterion['title']}{suffix}"
    if row_index % 9 == 0 and criterion["sc"] != "4.1.1":
        shaped += " Also applies to: EN 301 549 Criteria 9.1.1.1 (Web) Revised Section 508 501 (Web)(Software)"
    return shaped


def official_vpat_label(label: str, row_index: int) -> str:
    """Shape fixture labels like VPAT 2.5 WCAG criterion cells."""
    catalog = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
    sc = label.split(" ", 1)[0]
    criterion = next(item for item in catalog["criteria"] if item["sc"] == sc)
    return official_vpat_label_for_criterion(criterion, row_index)


def read_source() -> dict:
    source = json.loads(SOURCE_PATH.read_text(encoding="utf-8"))
    wcag = next(table for table in source["candidateTables"] if table["tableId"] == "table-wcag-22-a-aa")
    for row_index, row in enumerate(wcag["rows"], start=1):
        row["cells"][0] = official_vpat_label(row["cells"][0], row_index)
    return source


def build_union_rows() -> list[dict]:
    """Build the immutable 87-item WCAG version/level union evidence model."""
    catalog = json.loads(CATALOG_PATH.read_text(encoding="utf-8"))
    conformance_cycle = ("Supports", "Partially Supports", "Does Not Support", "Not Applicable")
    rows = []
    for row_index, criterion in enumerate(catalog["criteria"], start=1):
        conformance = "" if criterion["sc"] == "4.1.1" else conformance_cycle[(row_index - 1) % len(conformance_cycle)]
        if criterion["sc"] == UNION_CROSS_PAGE_SC:
            remarks = (
                "Synthetic cross-page evidence for SC 4.1.1 begins on one PDF page and continues after "
                "the repeated WCAG header on the next page. The continuation stays entirely within the "
                "criterion and remarks columns so the parser can preserve this literal evidence without "
                "guessing, while the DOCX and searchable PDF remain exact normalized-evidence peers."
            )
        elif row_index % 11 == 0:
            remarks = (
                f"Synthetic wrapped union evidence for SC {criterion['sc']} exercises a variable-height row "
                "with enough literal detail to wrap across multiple visual lines without changing the source evidence."
            )
        else:
            remarks = f"Synthetic union evidence for SC {criterion['sc']} across the declared WCAG scope."
        rows.append({
            "sourceRowIndex": row_index,
            "cells": [
                official_vpat_label_for_criterion(criterion, row_index),
                conformance,
                remarks,
            ],
        })
    if len(rows) != 87:
        raise ValueError("The full-union fixture must contain exactly 87 immutable catalog criteria.")
    return rows


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill.removeprefix("#"))


def set_cell_margins(cell, *, top: int = 70, start: int = 100, bottom: int = 70, end: int = 100) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for key, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_table_geometry(table, widths_twips: tuple[int, ...]) -> None:
    """Apply exact OOXML widths so LibreOffice/Word/Google Docs import agree."""
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths_twips)))

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_twips:
        grid_col = OxmlElement("w:gridCol")
        grid_col.set(qn("w:w"), str(width))
        grid.append(grid_col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            width = widths_twips[min(index, len(widths_twips) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(width))
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def style_table(table, widths_twips: tuple[int, ...]) -> None:
    set_table_geometry(table, widths_twips)
    header = table.rows[0]
    set_repeat_table_header(header)
    for cell in header.cells:
        set_cell_shading(cell, PURPLE)
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.color.rgb = RGBColor(255, 255, 255)
                run.font.name = "Arial"
                run.font.size = Pt(8)
    for row_index, row in enumerate(table.rows[1:], start=1):
        prevent_row_split(row)
        if row_index % 2 == 0:
            for cell in row.cells:
                set_cell_shading(cell, LIGHT_PURPLE)
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                paragraph.paragraph_format.line_spacing = 1.0
                for run in paragraph.runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(7.5)
                    run.font.color.rgb = RGBColor(31, 41, 55)


def add_source_table(document: Document, heading: str, table_model: dict, widths: tuple[int, ...]) -> None:
    document.add_heading(heading, level=2)
    headers = table_model["headers"]
    table = document.add_table(rows=1, cols=len(headers))
    for index, value in enumerate(headers):
        table.rows[0].cells[index].text = value
    for source_row in table_model["rows"]:
        row = table.add_row()
        for index, value in enumerate(source_row["cells"]):
            row.cells[index].text = value
    style_table(table, widths)


def configure_docx_styles(document: Document) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(9)
    normal.font.color.rgb = RGBColor(31, 41, 55)
    normal.paragraph_format.space_after = Pt(5)

    title = document.styles["Title"]
    title.font.name = "Arial"
    title.font.size = Pt(26)
    title.font.bold = False
    title.font.color.rgb = RGBColor(87, 6, 140)

    heading1 = document.styles["Heading 1"]
    heading1.font.name = "Arial"
    heading1.font.size = Pt(20)
    heading1.font.bold = True
    heading1.font.color.rgb = RGBColor(87, 6, 140)

    heading2 = document.styles["Heading 2"]
    heading2.font.name = "Arial"
    heading2.font.size = Pt(14)
    heading2.font.bold = True
    heading2.font.color.rgb = RGBColor(87, 6, 140)


def normalize_docx_zip(path: Path) -> None:
    with zipfile.ZipFile(path, "r") as source_zip:
        entries = [(info.filename, source_zip.read(info.filename)) for info in source_zip.infolist()]
    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temp_path = Path(handle.name)
    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target_zip:
            for filename, payload in sorted(entries):
                info = zipfile.ZipInfo(filename, (1980, 1, 1, 0, 0, 0))
                info.compress_type = zipfile.ZIP_DEFLATED
                info.create_system = 0
                info.external_attr = 0
                target_zip.writestr(info, payload, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
        temp_path.replace(path)
    finally:
        temp_path.unlink(missing_ok=True)


def build_docx(source: dict, path: Path, *, version: str = "2.5", include_wcag: bool = True) -> None:
    document = Document()
    configure_docx_styles(document)
    props = document.core_properties
    props.title = "Northstar Collaboration Suite Accessibility Conformance Report"
    props.subject = "Synthetic Stage 0 ingestion fixture"
    props.author = "NYU VPAT Analyzer synthetic fixture generator"
    props.keywords = "synthetic, VPAT, accessibility, test fixture"
    props.created = FIXED_TIME.replace(tzinfo=None)
    props.modified = FIXED_TIME.replace(tzinfo=None)
    props.last_modified_by = "NYU VPAT Analyzer synthetic fixture generator"
    props.revision = 1

    document.add_heading("Northstar Collaboration Suite", level=0)
    subtitle = document.add_paragraph()
    run = subtitle.add_run(f"Accessibility Conformance Report · VPAT {version}")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(87, 6, 140)
    for paragraph in source["bodyProse"]:
        document.add_paragraph(paragraph)

    metadata = source["candidateTables"][0]
    metadata_copy = json.loads(json.dumps(metadata))
    metadata_copy["rows"][1]["cells"][1] = version
    add_source_table(document, "Product and report details", metadata_copy, (2700, 6660))

    if include_wcag:
        wcag = source["candidateTables"][1]
        add_source_table(document, "WCAG 2.2 Level A and AA", wcag, (3600, 1800, 3960))

        duplicate = {
            "headers": wcag["headers"],
            "rows": [
                {
                    "sourceRowIndex": 1,
                    "cells": [
                        wcag["rows"][0]["cells"][0],
                        "Supports",
                        "Synthetic duplicate evidence retained for duplicate detection.",
                    ],
                }
            ],
        }
        add_source_table(document, "WCAG 2.2 duplicate-row test", duplicate, (3600, 1800, 3960))

    for table_model, heading, widths in (
        (source["candidateTables"][2], "Revised Section 508", (3600, 1800, 3960)),
        (source["candidateTables"][3], "EN 301 549", (3600, 1800, 3960)),
        (source["candidateTables"][4], "Internal release checklist", (5400, 3960)),
        (source["candidateTables"][5], "Literal source-text safety", (3600, 5760)),
    ):
        add_source_table(document, heading, table_model, widths)

    document.add_heading("Decorative merged layout table", level=2)
    merged = document.add_table(rows=2, cols=3)
    merged.rows[0].cells[0].merge(merged.rows[0].cells[2]).text = "Synthetic merged layout cell"
    merged.rows[1].cells[0].text = "Left"
    merged.rows[1].cells[1].text = "Middle"
    merged.rows[1].cells[2].text = "Right"
    style_table(merged, (3120, 3120, 3120))

    document.save(path)
    normalize_docx_zip(path)


def build_union_docx(path: Path, union_rows: list[dict]) -> None:
    """Create a faithful, synthetic VPAT DOCX spanning the full WCAG union."""
    document = Document()
    configure_docx_styles(document)
    props = document.core_properties
    props.title = "Northstar Full-Union Accessibility Conformance Report"
    props.subject = "Synthetic Stage 0 full WCAG union ingestion fixture"
    props.author = "NYU VPAT Analyzer synthetic fixture generator"
    props.keywords = "synthetic, VPAT, accessibility, WCAG union, test fixture"
    props.created = FIXED_TIME.replace(tzinfo=None)
    props.modified = FIXED_TIME.replace(tzinfo=None)
    props.last_modified_by = "NYU VPAT Analyzer synthetic fixture generator"
    props.revision = 1

    document.add_heading("Northstar Full-Union Accessibility Conformance Report", level=0)
    subtitle = document.add_paragraph()
    run = subtitle.add_run("Voluntary Product Accessibility Template® (VPAT®) · WCAG Edition · Version 2.5Rev")
    run.bold = True
    run.font.name = "Arial"
    run.font.size = Pt(11)
    run.font.color.rgb = RGBColor(87, 6, 140)
    document.add_paragraph(
        "This committed synthetic report declares WCAG 2.0, 2.1, and 2.2 at Levels A, AA, and AAA."
    )
    document.add_paragraph(
        "It contains all 87 immutable catalog criteria, including retired SC 4.1.1, and no production or user content."
    )

    metadata = {
        "headers": ["Field", "Value"],
        "rows": [
            {"sourceRowIndex": 1, "cells": ["Product name", "Northstar Full-Union Synthetic Fixture"]},
            {"sourceRowIndex": 2, "cells": ["VPAT template version", "2.5Rev"]},
            {"sourceRowIndex": 3, "cells": ["WCAG scope", "WCAG 2.0, 2.1, and 2.2; Levels A, AA, and AAA"]},
        ],
    }
    add_source_table(document, "Product and report details", metadata, (2700, 6660))
    add_source_table(
        document,
        "WCAG 2.0, 2.1, and 2.2 · Levels A, AA, and AAA",
        {
            "headers": ["Criteria", "Conformance Level", "Remarks and Explanations"],
            "rows": union_rows,
        },
        (3600, 1800, 3960),
    )
    document.save(path)
    normalize_docx_zip(path)


def build_xml_node_limit_docx(path: Path) -> None:
    """Create a small valid package that crosses the default XML-node boundary."""
    document = Document()
    props = document.core_properties
    props.title = "Northstar DOCX XML-node resource-limit fixture"
    props.subject = "Synthetic Stage 0 deterministic resource-limit fixture"
    props.author = "NYU VPAT Analyzer synthetic fixture generator"
    props.keywords = "synthetic, DOCX, XML node limit, test fixture"
    props.created = FIXED_TIME.replace(tzinfo=None)
    props.modified = FIXED_TIME.replace(tzinfo=None)
    props.last_modified_by = "NYU VPAT Analyzer synthetic fixture generator"
    props.revision = 1
    document.add_paragraph("Synthetic VPAT 2.5 DOCX resource-limit fixture; no production content.")
    document.save(path)

    with zipfile.ZipFile(path, "r") as source_zip:
        entries = {info.filename: source_zip.read(info.filename) for info in source_zip.infolist()}
    xml_names = [name for name in entries if name.lower().endswith((".xml", ".rels"))]
    base_node_count = sum(
        sum(1 for _ in ET.fromstring(entries[name]).iter())
        for name in xml_names
    )
    padding_count = DEFAULT_DOCX_MAX_XML_NODES + 32 - base_node_count
    if padding_count <= 0:
        raise ValueError("The base resource-limit DOCX unexpectedly exceeds its target XML-node count.")

    document_xml = entries["word/document.xml"].decode("utf-8")
    closing_body = "</w:body>"
    if document_xml.count(closing_body) != 1:
        raise ValueError("The resource-limit DOCX has an unexpected main-document structure.")
    document_xml = document_xml.replace(closing_body, "<w:p/>" * padding_count + closing_body)
    entries["word/document.xml"] = document_xml.encode("utf-8")

    with tempfile.NamedTemporaryFile(dir=path.parent, suffix=".docx", delete=False) as handle:
        temp_path = Path(handle.name)
    try:
        with zipfile.ZipFile(temp_path, "w", compression=zipfile.ZIP_DEFLATED, compresslevel=9) as target_zip:
            for filename, payload in sorted(entries.items()):
                info = zipfile.ZipInfo(filename, (1980, 1, 1, 0, 0, 0))
                info.create_system = 0
                info.external_attr = 0
                if filename == "word/document.xml":
                    info.compress_type = zipfile.ZIP_STORED
                    target_zip.writestr(info, payload, compress_type=zipfile.ZIP_STORED)
                else:
                    info.compress_type = zipfile.ZIP_DEFLATED
                    target_zip.writestr(info, payload, compress_type=zipfile.ZIP_DEFLATED, compresslevel=9)
        temp_path.replace(path)
    finally:
        temp_path.unlink(missing_ok=True)

    if path.stat().st_size >= DEFAULT_DOCX_MAX_SOURCE_BYTES:
        raise ValueError("The resource-limit DOCX must remain below the default source-byte cap.")


def draw_pdf_header(pdf: canvas.Canvas, subtitle: str, page_number: int) -> float:
    width, height = landscape(letter)
    pdf.setFillColor(HexColor(PURPLE))
    pdf.setFont("Helvetica-Bold", 17)
    pdf.drawString(28, height - 31, "Northstar Collaboration Suite")
    pdf.setFillColor(HexColor(INK))
    pdf.setFont("Helvetica", 8.5)
    pdf.drawRightString(width - 28, height - 29, f"Synthetic VPAT 2.5 · page {page_number}")
    pdf.setFont("Helvetica-Bold", 10)
    pdf.drawString(28, height - 49, subtitle)
    return height - 70


def fitted_font_size(text: str, max_width: float, base_size: float = 6.6, minimum: float = 4.8) -> float:
    size = base_size
    while size > minimum and pdfmetrics.stringWidth(text, "Helvetica", size) > max_width:
        size -= 0.2
    return max(size, minimum)


def draw_pdf_table(pdf: canvas.Canvas, y: float, headers: list[str], rows: list[dict], widths: tuple[float, ...]) -> float:
    left = 28.0
    row_height = 15.0
    total_width = sum(widths)
    pdf.setFillColor(HexColor(PURPLE))
    pdf.rect(left, y - row_height + 3, total_width, row_height, fill=1, stroke=0)
    x = left
    for index, header in enumerate(headers):
        pdf.setFillColor(white)
        pdf.setFont("Helvetica-Bold", 6.8)
        pdf.drawString(x + 3, y - 7, header)
        x += widths[index]
    y -= row_height
    for row_index, source_row in enumerate(rows):
        if row_index % 2 == 1:
            pdf.setFillColor(HexColor(LIGHT_PURPLE))
            pdf.rect(left, y - row_height + 3, total_width, row_height, fill=1, stroke=0)
        x = left
        for index, cell in enumerate(source_row["cells"]):
            pdf.setFillColor(HexColor(INK))
            size = fitted_font_size(str(cell), widths[index] - 6)
            pdf.setFont("Helvetica", size)
            pdf.drawString(x + 3, y - 7, str(cell))
            x += widths[index]
        y -= row_height
    return y


def wrap_pdf_cell(text: str, max_width: float, *, font: str = "Helvetica", size: float = 6.2) -> list[str]:
    """Greedily wrap a cell while preserving its normalized word sequence."""
    words = str(text).split()
    if not words:
        return []
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        proposed = f"{current} {word}"
        if pdfmetrics.stringWidth(proposed, font, size) <= max_width:
            current = proposed
        else:
            lines.append(current)
            current = word
    lines.append(current)
    if any(pdfmetrics.stringWidth(line, font, size) > max_width for line in lines):
        raise ValueError("A full-union PDF cell contains an unbreakable over-width token.")
    return lines


def union_pdf_lines(row: dict, widths: tuple[float, ...], font_size: float) -> list[list[str]]:
    return [
        wrap_pdf_cell(cell, widths[index] - 8, size=font_size)
        for index, cell in enumerate(row["cells"])
    ]


def draw_union_pdf_header_row(
    pdf: canvas.Canvas,
    y: float,
    headers: tuple[str, ...],
    widths: tuple[float, ...],
) -> float:
    left = 28.0
    height = 18.0
    pdf.setFillColor(HexColor(PURPLE))
    pdf.rect(left, y - height, sum(widths), height, fill=1, stroke=0)
    x = left
    for index, header in enumerate(headers):
        pdf.setFillColor(white)
        pdf.setFont("Helvetica-Bold", 6.8)
        pdf.drawString(x + 4, y - 11, header)
        x += widths[index]
    return y - height


def draw_union_pdf_row_segment(
    pdf: canvas.Canvas,
    y: float,
    lines_by_cell: list[list[str]],
    widths: tuple[float, ...],
    row_index: int,
    *,
    font_size: float = 6.2,
    line_height: float = 8.2,
) -> float:
    left = 28.0
    line_count = max((len(lines) for lines in lines_by_cell), default=1)
    height = max(15.0, line_count * line_height + 5.0)
    if row_index % 2 == 0:
        pdf.setFillColor(HexColor(LIGHT_PURPLE))
    else:
        pdf.setFillColor(white)
    pdf.rect(left, y - height, sum(widths), height, fill=1, stroke=0)
    x = left
    for cell_index, lines in enumerate(lines_by_cell):
        pdf.setFillColor(HexColor(INK))
        pdf.setFont("Helvetica", font_size)
        for line_index, line in enumerate(lines):
            pdf.drawString(x + 4, y - 8.0 - line_index * line_height, line)
        x += widths[cell_index]
    pdf.setStrokeColor(Color(87 / 255, 6 / 255, 140 / 255, alpha=0.18))
    pdf.setLineWidth(0.35)
    pdf.line(left, y - height, left + sum(widths), y - height)
    return y - height


def start_union_pdf_page(pdf: canvas.Canvas, page_number: int) -> float:
    scope = "WCAG 2.0, 2.1, and 2.2 - Levels A, AA, and AAA"
    y = draw_pdf_header(pdf, scope, page_number)
    pdf.setFillColor(HexColor(INK))
    pdf.setFont("Helvetica-Bold", 8.5)
    pdf.drawString(28, y, scope)
    y -= 11
    return draw_union_pdf_header_row(
        pdf,
        y,
        ("Criteria", "Conformance Level", "Remarks and Explanations"),
        (326.0, 106.0, 304.0),
    )


def build_union_searchable_pdf(path: Path, union_rows: list[dict]) -> None:
    """Create a searchable, variable-height PDF with one safe page-spanning row."""
    pdf = canvas.Canvas(str(path), pagesize=landscape(letter), pageCompression=1, invariant=1)
    pdf.setTitle("Northstar Full-Union Accessibility Conformance Report")
    pdf.setAuthor("NYU VPAT Analyzer synthetic fixture generator")
    pdf.setSubject("Synthetic Stage 0 full WCAG union searchable PDF fixture")
    widths = (326.0, 106.0, 304.0)
    font_size = 6.2
    bottom = 30.0
    page_number = 1
    y = start_union_pdf_page(pdf, page_number)
    split_exercised = False

    for row_index, row in enumerate(union_rows, start=1):
        lines_by_cell = union_pdf_lines(row, widths, font_size)
        line_count = max(len(lines) for lines in lines_by_cell)
        row_height = max(15.0, line_count * 8.2 + 5.0)

        if row["cells"][0].startswith(f"{UNION_CROSS_PAGE_SC} "):
            if y - 15.0 < bottom:
                pdf.showPage()
                page_number += 1
                y = start_union_pdf_page(pdf, page_number)
            first_segment = [lines[:1] for lines in lines_by_cell]
            continuation = [lines[1:] for lines in lines_by_cell]
            if not continuation[2]:
                raise ValueError("The designated PDF cross-page row must wrap in its remarks cell.")
            y = draw_union_pdf_row_segment(pdf, y, first_segment, widths, row_index)
            pdf.showPage()
            page_number += 1
            y = start_union_pdf_page(pdf, page_number)
            continuation_height = max(15.0, max(len(lines) for lines in continuation) * 8.2 + 5.0)
            if y - continuation_height < bottom:
                raise ValueError("The designated PDF continuation does not fit on a fresh page.")
            y = draw_union_pdf_row_segment(pdf, y, continuation, widths, row_index)
            split_exercised = True
            continue

        if y - row_height < bottom:
            pdf.showPage()
            page_number += 1
            y = start_union_pdf_page(pdf, page_number)
        y = draw_union_pdf_row_segment(pdf, y, lines_by_cell, widths, row_index)

    if not split_exercised:
        raise ValueError("The full-union PDF did not exercise its cross-page continuation row.")
    pdf.save()


def build_searchable_pdf(source: dict, path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=landscape(letter), pageCompression=1, invariant=1)
    pdf.setTitle("Northstar Collaboration Suite Accessibility Conformance Report")
    pdf.setAuthor("NYU VPAT Analyzer synthetic fixture generator")
    pdf.setSubject("Synthetic Stage 0 searchable PDF fixture")
    page_number = 1

    y = draw_pdf_header(pdf, "Product details and WCAG 2.2 A/AA · part 1", page_number)
    for paragraph in source["bodyProse"]:
        pdf.setFont("Helvetica", 7.5)
        pdf.setFillColor(HexColor(INK))
        pdf.drawString(28, y, paragraph)
        y -= 12
    y -= 3
    y = draw_pdf_table(pdf, y, source["candidateTables"][0]["headers"], source["candidateTables"][0]["rows"], (180, 556))
    y -= 10
    wcag = source["candidateTables"][1]
    first_half = wcag["rows"][:25]
    pdf.setFont("Helvetica-Bold", 8.5)
    pdf.drawString(28, y, "WCAG 2.2 Level A and AA")
    y -= 10
    draw_pdf_table(pdf, y, wcag["headers"], first_half, (326, 106, 304))
    pdf.showPage()

    page_number += 1
    y = draw_pdf_header(pdf, "WCAG 2.2 A/AA · part 2", page_number)
    pdf.setFont("Helvetica-Bold", 8.5)
    pdf.drawString(28, y, "WCAG 2.2 Level A and AA")
    y -= 10
    draw_pdf_table(pdf, y, wcag["headers"], wcag["rows"][25:], (326, 106, 304))
    pdf.showPage()

    page_number += 1
    y = draw_pdf_header(pdf, "Exclusion and duplicate cases", page_number)
    duplicate = [{"sourceRowIndex": 1, "cells": [wcag["rows"][0]["cells"][0], "Supports", "Synthetic duplicate evidence retained for duplicate detection."]}]
    pdf.setFont("Helvetica-Bold", 8.5)
    pdf.drawString(28, y, "WCAG 2.2 duplicate-row test")
    y -= 10
    y = draw_pdf_table(pdf, y, wcag["headers"], duplicate, (326, 106, 304))
    y -= 10
    for table_model, heading, widths in (
        (source["candidateTables"][2], "Revised Section 508", (326, 106, 304)),
        (source["candidateTables"][3], "EN 301 549", (326, 106, 304)),
        (source["candidateTables"][4], "Internal release checklist", (430, 306)),
        (source["candidateTables"][5], "Literal source-text safety", (300, 436)),
    ):
        pdf.setFont("Helvetica-Bold", 8.5)
        pdf.drawString(28, y, heading)
        y -= 10
        y = draw_pdf_table(pdf, y, table_model["headers"], table_model["rows"], widths)
        y -= 10
    pdf.save()


def build_image_only_pdf(source: dict, path: Path) -> None:
    width, height = landscape(letter)
    scale = 2
    image = Image.new("RGB", (int(width * scale), int(height * scale)), "white")
    draw = ImageDraw.Draw(image)
    try:
        font_title = ImageFont.truetype("Arial.ttf", 32)
        font_body = ImageFont.truetype("Arial.ttf", 16)
    except OSError:
        font_title = ImageFont.load_default()
        font_body = ImageFont.load_default()
    draw.text((56, 50), "Northstar Collaboration Suite", fill=(87, 6, 140), font=font_title)
    draw.text((56, 105), "Synthetic image-only VPAT 2.5 page — OCR is intentionally forbidden.", fill=(31, 41, 55), font=font_body)
    y = 155
    for row in source["candidateTables"][1]["rows"][:12]:
        draw.text((56, y), " | ".join(row["cells"]), fill=(31, 41, 55), font=font_body)
        y += 34
    buffer = io.BytesIO()
    image.save(buffer, format="PNG", optimize=False, compress_level=9)
    buffer.seek(0)
    pdf = canvas.Canvas(str(path), pagesize=landscape(letter), pageCompression=1, invariant=1)
    pdf.setTitle("Northstar image-only synthetic fixture")
    pdf.drawImage(ImageReader(buffer), 0, 0, width=width, height=height, preserveAspectRatio=False, mask="auto")
    pdf.save()


def build_insufficient_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter, pageCompression=1, invariant=1)
    pdf.setTitle("Northstar insufficient-text synthetic fixture")
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(54, 730, "Northstar accessibility note")
    pdf.setFont("Helvetica", 10)
    pdf.drawString(54, 708, "Short searchable text; no recoverable VPAT table is present.")
    pdf.save()


def build_ambiguous_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=landscape(letter), pageCompression=1, invariant=1)
    y = draw_pdf_header(pdf, "Ambiguous WCAG row rejection case", 1)
    pdf.setFont("Helvetica", 7.5)
    pdf.drawString(28, y, "This synthetic file contains enough searchable text but deliberately combines multiple success criteria in one row.")
    y -= 18
    rows = [
        {"sourceRowIndex": 1, "cells": ["1.1.1 / 1.2.1 conflicting criterion", "Supports", "Synthetic ambiguity; the parser must reject rather than guess."]},
        {"sourceRowIndex": 2, "cells": ["2.1.1 Keyboard", "Supports", "Additional searchable synthetic evidence for deterministic table detection."]},
    ]
    draw_pdf_table(pdf, y, ["Criteria", "Conformance Level", "Remarks and Explanations"], rows, (326, 106, 304))
    pdf.save()


def build_resource_limit_pdf(path: Path) -> None:
    pdf = canvas.Canvas(str(path), pagesize=letter, pageCompression=1, invariant=1)
    for page_number in range(1, 10):
        pdf.setFont("Helvetica-Bold", 14)
        pdf.drawString(54, 730, f"Northstar resource-limit fixture page {page_number}")
        pdf.setFont("Helvetica", 9)
        for index in range(40):
            pdf.drawString(54, 700 - index * 15, f"Synthetic bounded parser work item {page_number:02d}-{index:02d}; no production content.")
        pdf.showPage()
    pdf.save()


def build_encrypted_pdf(source_pdf: Path, path: Path) -> None:
    reader = PdfReader(str(source_pdf))
    writer = PdfWriter()
    for page in reader.pages:
        writer.add_page(page)
    writer.add_metadata({
        "/Title": "Northstar encrypted synthetic fixture",
        "/Author": "NYU VPAT Analyzer synthetic fixture generator",
        "/CreationDate": "D:20260812120000Z",
        "/ModDate": "D:20260812120000Z",
    })
    fixed_id = hashlib.sha256(b"northstar-encrypted-stage0-fixture").digest()[:16]
    writer._ID = ArrayObject((ByteStringObject(fixed_id), ByteStringObject(fixed_id)))
    writer.encrypt(FIXTURE_PASSWORD, owner_password=FIXTURE_PASSWORD, algorithm="RC4-128")
    with path.open("wb") as handle:
        writer.write(handle)


def build_malformed_pdf(valid_pdf: Path, path: Path) -> None:
    payload = valid_pdf.read_bytes()
    # Keep the header and an incomplete object graph, but remove xref/trailer.
    path.write_bytes(payload[: min(900, len(payload) // 5)])


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def write_manifest(files: list[dict]) -> None:
    manifest = {
        "id": "stage0-real-format-fixtures",
        "version": "1.0.0",
        "generatorVersion": GENERATOR_VERSION,
        "sourceFixture": str(SOURCE_PATH.relative_to(ROOT)),
        "sourceSha256": sha256(SOURCE_PATH),
        "catalogSource": str(CATALOG_PATH.relative_to(ROOT)),
        "catalogSha256": sha256(CATALOG_PATH),
        "designPreset": PARSER_DESIGN_PRESET,
        "syntheticOnly": True,
        "containsProductionContent": False,
        "files": [],
    }
    for item in files:
        path = OUTPUT_DIR / item["filename"]
        manifest["files"].append({
            **item,
            "bytes": path.stat().st_size,
            "sha256": sha256(path),
        })
    manifest_path = OUTPUT_DIR / "manifest.v1.json"
    manifest_path.write_text(json.dumps(manifest, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def generate() -> None:
    source = read_source()
    union_rows = build_union_rows()
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    files = [
        {"id": "primary-docx", "filename": "northstar-vpat-2.5.docx", "sourceType": "docx", "expectedStatus": "complete", "expectedUniqueCriteria": 55, "expectedDuplicateGroups": 1},
        {"id": "primary-searchable-pdf", "filename": "northstar-vpat-2.5-searchable.pdf", "sourceType": "pdf", "expectedStatus": "complete", "expectedUniqueCriteria": 55, "expectedDuplicateGroups": 1},
        {"id": "scope-union-docx", "filename": "northstar-vpat-2.5-scope-union.docx", "sourceType": "docx", "expectedStatus": "complete", "expectedUniqueCriteria": 87, "expectedMissingCriteria": 0, "expectedDuplicateGroups": 0},
        {"id": "scope-union-pdf", "filename": "northstar-vpat-2.5-scope-union.pdf", "sourceType": "pdf", "expectedStatus": "complete", "expectedUniqueCriteria": 87, "expectedMissingCriteria": 0, "expectedDuplicateGroups": 0},
        {"id": "image-only-pdf", "filename": "northstar-vpat-2.5-image-only.pdf", "sourceType": "pdf", "expectedStatus": "rejected", "expectedCode": "PDF_NON_SEARCHABLE"},
        {"id": "insufficient-text-pdf", "filename": "northstar-vpat-2.5-insufficient-text.pdf", "sourceType": "pdf", "expectedStatus": "rejected", "expectedCode": "PDF_INSUFFICIENT_TEXT"},
        {"id": "malformed-pdf", "filename": "northstar-vpat-2.5-malformed.pdf", "sourceType": "pdf", "expectedStatus": "rejected", "expectedCode": "SOURCE_MALFORMED"},
        {"id": "encrypted-pdf", "filename": "northstar-vpat-2.5-encrypted.pdf", "sourceType": "pdf", "expectedStatus": "rejected", "expectedCode": "PDF_ENCRYPTED"},
        {"id": "ambiguous-pdf", "filename": "northstar-vpat-2.5-ambiguous.pdf", "sourceType": "pdf", "expectedStatus": "rejected", "expectedCode": "WCAG_ROWS_AMBIGUOUS"},
        {"id": "resource-limit-pdf", "filename": "northstar-vpat-2.5-resource-limit.pdf", "sourceType": "pdf", "expectedStatus": "rejected", "expectedCode": "RESOURCE_LIMIT_EXCEEDED", "testLimit": "maxPdfPages=2"},
        {"id": "resource-limit-docx", "filename": "northstar-vpat-2.5-resource-limit.docx", "sourceType": "docx", "expectedStatus": "rejected", "expectedCode": "RESOURCE_LIMIT_EXCEEDED", "resourceBoundary": "default-maxXmlNodes=200000"},
        {"id": "unsupported-version-docx", "filename": "northstar-vpat-2.4.docx", "sourceType": "docx", "expectedStatus": "rejected", "expectedCode": "VPAT_VERSION_UNSUPPORTED"},
        {"id": "no-wcag-docx", "filename": "northstar-vpat-2.5-no-wcag.docx", "sourceType": "docx", "expectedStatus": "rejected", "expectedCode": "WCAG_TABLE_NOT_FOUND"},
        {"id": "unsupported-text", "filename": "northstar-vpat.txt", "sourceType": "unknown", "expectedStatus": "rejected", "expectedCode": "SOURCE_TYPE_UNSUPPORTED"},
    ]

    build_docx(source, OUTPUT_DIR / "northstar-vpat-2.5.docx")
    build_searchable_pdf(source, OUTPUT_DIR / "northstar-vpat-2.5-searchable.pdf")
    build_union_docx(OUTPUT_DIR / "northstar-vpat-2.5-scope-union.docx", union_rows)
    build_union_searchable_pdf(OUTPUT_DIR / "northstar-vpat-2.5-scope-union.pdf", union_rows)
    build_image_only_pdf(source, OUTPUT_DIR / "northstar-vpat-2.5-image-only.pdf")
    build_insufficient_pdf(OUTPUT_DIR / "northstar-vpat-2.5-insufficient-text.pdf")
    build_ambiguous_pdf(OUTPUT_DIR / "northstar-vpat-2.5-ambiguous.pdf")
    build_resource_limit_pdf(OUTPUT_DIR / "northstar-vpat-2.5-resource-limit.pdf")
    build_xml_node_limit_docx(OUTPUT_DIR / "northstar-vpat-2.5-resource-limit.docx")
    build_encrypted_pdf(OUTPUT_DIR / "northstar-vpat-2.5-searchable.pdf", OUTPUT_DIR / "northstar-vpat-2.5-encrypted.pdf")
    build_malformed_pdf(OUTPUT_DIR / "northstar-vpat-2.5-searchable.pdf", OUTPUT_DIR / "northstar-vpat-2.5-malformed.pdf")
    build_docx(source, OUTPUT_DIR / "northstar-vpat-2.4.docx", version="2.4")
    build_docx(source, OUTPUT_DIR / "northstar-vpat-2.5-no-wcag.docx", include_wcag=False)
    (OUTPUT_DIR / "northstar-vpat.txt").write_text("Synthetic Northstar unsupported source type.\n", encoding="utf-8")
    write_manifest(files)


def verify() -> None:
    manifest_path = OUTPUT_DIR / "manifest.v1.json"
    if not manifest_path.exists():
        raise SystemExit("fixture manifest is missing; run without --verify first")
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    failures = []
    if manifest.get("sourceSha256") != sha256(SOURCE_PATH):
        failures.append("source fixture hash differs")
    if manifest.get("catalogSha256") != sha256(CATALOG_PATH):
        failures.append("catalog hash differs")
    for item in manifest.get("files", []):
        path = OUTPUT_DIR / item["filename"]
        if not path.exists():
            failures.append(f"missing {item['filename']}")
            continue
        if path.stat().st_size != item["bytes"]:
            failures.append(f"byte count differs for {item['filename']}")
        if sha256(path) != item["sha256"]:
            failures.append(f"SHA-256 differs for {item['filename']}")
    if failures:
        raise SystemExit("fixture verification failed:\n- " + "\n- ".join(failures))
    print(f"verified {len(manifest['files'])} synthetic real-format fixtures")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--verify", action="store_true", help="verify committed files against the manifest")
    args = parser.parse_args()
    if args.verify:
        verify()
    else:
        generate()
        verify()


if __name__ == "__main__":
    main()
